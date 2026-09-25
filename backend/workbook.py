from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties
from pydantic import ValidationError

from .engine import calculate
from .models import AgentRow, Estimate

IMPORT_COLUMNS = [
    "name",
    "complexity",
    "count",
    "invocations",
    "model_id",
    "calls",
    "input_tokens",
    "output_tokens",
    "retry_rate",
    "cache_fraction",
    "cache_write_fraction",
]


def literal(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def finish(workbook):
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="153C39")
            cell.alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                48, max(16, len(str(column[0].value or "")) + 3)
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)) or cell.data_type == "f":
                    cell.number_format = "#,##0.000000;[Red]-#,##0.000000"
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Agents"
    literal(ws, IMPORT_COLUMNS)
    literal(ws, ["Document classification", "simple", 1, 1000, "", "", "", "", "", "", ""])
    notes = wb.create_sheet("Instructions")
    literal(notes, ["Field", "Meaning"])
    for key, value in [
        ("Required", "name, complexity, count, invocations; complexity is simple, medium, or high."),
        ("model_id", "Exact catalog or custom model ID. Blank inherits the profile model."),
        (
            "Units",
            "invocations per agent/month; calls per invocation; tokens per call; rates are fractions, e.g. 0.02.",
        ),
        ("Overrides", "Blank execution cells inherit the profile. Zero is an explicit override."),
        (
            "Ownership",
            "Each row is a disjoint group. This import replaces the current agent list after preview.",
        ),
        ("Safety", "Values only: formulas, macros and external links are not supported. Maximum 1,000 rows."),
    ]:
        literal(notes, [key, value])
    return finish(wb)


def read_import(data: bytes):
    errors, agents = [], []
    if len(data) > 5_000_000:
        raise ValueError("Spreadsheet must be smaller than 5 MB.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 30_000_000:
                raise ValueError("Spreadsheet expands beyond the 30 MB limit.")
            if any(
                "vbaproject" in f.filename.lower() or "externallinks/" in f.filename.lower()
                for f in archive.infolist()
            ):
                raise ValueError("Macros and external workbook links are not accepted.")
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Upload a valid .xlsx workbook.") from exc
    try:
        ws = wb["Agents"] if "Agents" in wb.sheetnames else wb.active
        if ws.max_row > 1001 or ws.max_column > 30:
            raise ValueError("Use at most 1,000 agent rows and 30 columns.")
        headers = [str(c.value or "").strip() for c in next(ws.iter_rows())]
        if len(set(headers)) != len(headers) or set(headers) - set(IMPORT_COLUMNS):
            raise ValueError("Use unique column names from the provided import template.")
        if not {"name", "complexity", "count", "invocations"}.issubset(headers):
            raise ValueError("Required columns: name, complexity, count, invocations.")
        for index, cells in enumerate(ws.iter_rows(min_row=2), 2):
            if all(c.value is None for c in cells):
                continue
            if any(c.data_type == "f" for c in cells):
                errors.append(f"Row {index}: formulas are not allowed; paste values instead.")
                continue
            values = {key: c.value for key, c in zip(headers, cells) if c.value is not None and c.value != ""}
            missing = {"name", "complexity", "count", "invocations"} - values.keys()
            if missing:
                errors.append(f"Row {index}: required values missing: {', '.join(sorted(missing))}.")
                continue
            overrides = {k: values.pop(k) for k in list(values) if k in IMPORT_COLUMNS[4:]}
            try:
                agents.append(AgentRow(**values, overrides=overrides).model_dump(mode="json"))
            except ValidationError as exc:
                errors.extend(
                    f"Row {index}, {'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors()
                )
        if not agents and not errors:
            errors.append("No agent rows found.")
    finally:
        wb.close()
    return {"agents": agents if not errors else [], "errors": errors}


def export_estimate(estimate: Estimate):
    result = calculate(estimate)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    literal(
        summary,
        [
            "Scenario",
            "LLM monthly USD",
            "Additional monthly USD",
            "Recurring monthly USD",
            "One-time USD",
            "First month USD",
            "First year USD",
            "Status",
        ],
    )
    calc = wb.create_sheet("Calculations")
    columns = [
        "Scenario",
        "Agent/group",
        "Step",
        "Model",
        "Count",
        "Invocations/agent/month",
        "Calls/invocation",
        "Extra attempt rate",
        "Input/call",
        "Output/call",
        "Cache read fraction",
        "Cache write fraction",
        "Input USD/M",
        "Output USD/M",
        "Cache read USD/M",
        "Cache write USD/M",
        "Monthly calls",
        "Input tokens/month",
        "Output tokens/month",
        "Input USD",
        "Output USD",
        "Cache read USD",
        "Cache write USD",
        "LLM monthly USD",
        "Status",
    ]
    literal(calc, columns)
    for s in result["scenarios"]:
        start = calc.max_row + 1
        for line in s["lines"]:
            rates = line["rates"]
            literal(
                calc,
                [
                    s["name"],
                    line["name"],
                    line["step"],
                    line["model_id"],
                    line["count"],
                    *[
                        float(line[k])
                        for k in (
                            "invocations",
                            "calls_per_invocation",
                            "retry_rate",
                            "input_per_call",
                            "output_per_call",
                            "cache_fraction",
                            "cache_write_fraction",
                        )
                    ],
                    *[
                        float(rates[k]) if rates[k] is not None else None
                        for k in ("input", "output", "cache_read", "cache_write")
                    ],
                    *[None] * 8,
                    "; ".join(line["issues"]) or "Complete",
                ],
            )
            r = calc.max_row
            formulas = {"Q": f"E{r}*F{r}*G{r}*(1+H{r})", "R": f"Q{r}*I{r}", "S": f"Q{r}*J{r}"}
            if not line["issues"]:
                formulas |= {
                    "T": f"R{r}*(1-K{r}-L{r})*M{r}/1000000",
                    "U": f"S{r}*N{r}/1000000",
                    "V": f"R{r}*K{r}*O{r}/1000000",
                    "W": f"R{r}*L{r}*P{r}/1000000",
                    "X": f"SUM(T{r}:W{r})",
                }
            for col, formula in formulas.items():
                calc[f"{col}{r}"] = "=" + formula
        end = calc.max_row
        literal(
            summary,
            [
                s["name"],
                None,
                float(result["recurring"]),
                None,
                float(result["one_time"]),
                None,
                None,
                "Complete" if s["complete"] else "INCOMPLETE — known costs only",
            ],
        )
        r = summary.max_row
        summary[f"B{r}"] = f"=SUM(Calculations!X{start}:X{end})" if end >= start else "=0"
        summary[f"D{r}"] = f"=B{r}+C{r}"
        summary[f"F{r}"] = f"=D{r}+E{r}"
        summary[f"G{r}"] = f"=D{r}*12+E{r}"

    agents = wb.create_sheet("Agents")
    literal(agents, IMPORT_COLUMNS)
    for row in estimate.agents:
        literal(
            agents,
            [
                row.name,
                row.complexity,
                row.count,
                float(row.invocations),
                *[
                    str(getattr(row.overrides, k)) if getattr(row.overrides, k) is not None else ""
                    for k in IMPORT_COLUMNS[4:]
                ],
            ],
        )
    for title, records in [
        ("Profiles", [dict(complexity=k, **v.model_dump(mode="json")) for k, v in estimate.profiles.items()]),
        ("Scenarios", [s.model_dump(mode="json") for s in estimate.scenarios]),
        ("Additional costs", [c.model_dump(mode="json") for c in estimate.additional_costs]),
    ]:
        sheet = wb.create_sheet(title)
        if records:
            keys = list(records[0])
            literal(sheet, keys)
            for record in records:
                literal(sheet, [record[k] for k in keys])
        else:
            literal(sheet, ["No items"])
    pricing = wb.create_sheet("Pricing")
    literal(
        pricing,
        [
            "Model ID",
            "Provider",
            "Input USD/M",
            "Output USD/M",
            "Cache read USD/M",
            "Cache write USD/M",
            "Source",
            "Retrieved at",
            "Custom",
            "Tier details",
            "Unsupported",
        ],
    )
    for p in estimate.prices.values():
        literal(
            pricing,
            [
                p.id,
                p.provider,
                *[
                    float(getattr(p, k)) if getattr(p, k) is not None else None
                    for k in ("input", "output", "cache_read", "cache_write")
                ],
                p.source,
                p.retrieved_at,
                p.custom,
                "; ".join(t.model_dump_json() for t in p.tiers),
                ", ".join(p.unsupported),
            ],
        )
    notes = wb.create_sheet("Read me")
    literal(notes, ["Item", "Detail"])
    for key, value in [
        ("Estimate", estimate.name),
        ("Notes", estimate.notes),
        ("Currency", "USD"),
        (
            "Calculation editing",
            "Edit numeric inputs in Calculations to recalculate costs. Profiles, Scenarios, Agents and Pricing document the snapshot; editing those sheets does not propagate to Calculations.",
        ),
        (
            "Tier rates",
            "Rates in Calculations are selected using the exported per-call input size. After crossing a tier threshold, update the applicable rates from Pricing or re-export from the app.",
        ),
        (
            "Partial totals",
            "Incomplete rows have blank costs; summaries are known subtotals, not full budgets.",
        ),
        (
            "Annualization",
            "First year = 12 identical recurring months + one-time items. Additional monthly/one-time values in Summary are editable frozen inputs.",
        ),
        (
            "Scope",
            "Text tokens, including total billable reasoning in output. Cache writes use supplied rates. Tools, cache storage, multimodal charges require additional items/custom rates.",
        ),
        (
            "Scenarios",
            "Planning assumptions, not statistical confidence intervals. Execution fields in Calculations are scenario-effective values.",
        ),
        (
            "Detailed workflows",
            "Each step appears in Calculations and replaces the aggregate. The Agents sheet is an aggregate import template, not a lossless backup of detailed steps.",
        ),
        ("Version", f"Schema {estimate.schema_version}; defaults {estimate.defaults_version}"),
    ]:
        literal(notes, [key, value])
    return finish(wb)
